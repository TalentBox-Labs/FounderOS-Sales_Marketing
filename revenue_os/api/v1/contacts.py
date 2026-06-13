from __future__ import annotations

import uuid
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from pydantic import BaseModel, EmailStr
from sqlalchemy.orm import Session

from revenue_os.database import get_db
from revenue_os.integrations.n8n import new_lead_webhook
from revenue_os.models.contact import Company, Contact, ContactSource, ContactStatus
from revenue_os.services.scoring_service import score_contact
from revenue_os.services.search_service import delete_index, index_contact, search

router = APIRouter(prefix="/contacts", tags=["contacts"])


class ContactCreate(BaseModel):
    first_name: str
    last_name: str
    company_id: Optional[str] = None
    designation: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    linkedin_url: Optional[str] = None
    twitter_url: Optional[str] = None
    source: ContactSource = ContactSource.MANUAL
    status: ContactStatus = ContactStatus.LEAD
    tags: Optional[str] = None
    notes: Optional[str] = None


class ContactUpdate(BaseModel):
    first_name: Optional[str] = None
    last_name: Optional[str] = None
    company_id: Optional[str] = None
    designation: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    linkedin_url: Optional[str] = None
    twitter_url: Optional[str] = None
    status: Optional[ContactStatus] = None
    tags: Optional[str] = None
    notes: Optional[str] = None


class ContactResponse(BaseModel):
    id: uuid.UUID
    first_name: str
    last_name: str
    full_name: str
    company_id: Optional[uuid.UUID] = None
    designation: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    linkedin_url: Optional[str] = None
    twitter_url: Optional[str] = None
    source: ContactSource
    status: ContactStatus
    lead_score: int
    tags: Optional[str] = None
    notes: Optional[str] = None
    created_at: datetime
    updated_at: datetime

    model_config = {"from_attributes": True}


@router.get("", response_model=list[ContactResponse])
def list_contacts(
    status: Optional[ContactStatus] = Query(None),
    source: Optional[ContactSource] = Query(None),
    search: Optional[str] = Query(None),
    skip: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=1000),
    db: Session = Depends(get_db),
):
    query = db.query(Contact)
    if status:
        query = query.filter(Contact.status == status)
    if source:
        query = query.filter(Contact.source == source)
    if search:
        pattern = f"%{search}%"
        query = query.filter(
            Contact.first_name.ilike(pattern)
            | Contact.last_name.ilike(pattern)
            | Contact.email.ilike(pattern)
            | Contact.designation.ilike(pattern)
        )
    return query.offset(skip).limit(limit).all()


@router.get("/{contact_id}/timeline")
def contact_timeline(
    contact_id: str,
    limit: int = Query(50, ge=1, le=500),
    db: Session = Depends(get_db),
):
    from revenue_os.models.activity import Activity
    from revenue_os.models.activity import EmailActivity

    contact = db.query(Contact).filter(Contact.id == contact_id).first()
    if not contact:
        raise HTTPException(status_code=404, detail="Contact not found")

    activities = (
        db.query(Activity)
        .filter(Activity.contact_id == contact_id)
        .order_by(Activity.performed_at.desc())
        .limit(limit)
        .all()
    )

    results = []
    for a in activities:
        item = {
            "id": str(a.id),
            "type": a.activity_type.value,
            "subject": a.subject,
            "body": a.body[:500] if a.body else None,
            "direction": a.direction,
            "status": a.status,
            "performed_at": a.performed_at.isoformat() if a.performed_at else None,
        }
        if a.activity_type.value in ("email", "email_open", "email_click", "email_reply"):
            ea = (
                db.query(EmailActivity)
                .filter(EmailActivity.activity_id == a.id)
                .first()
            )
            if ea:
                item["email"] = {
                    "message_id": ea.message_id,
                    "from": ea.from_address,
                    "to": ea.to_addresses,
                    "opened_at": ea.opened_at.isoformat() if ea.opened_at else None,
                    "clicked_at": ea.clicked_at.isoformat() if ea.clicked_at else None,
                }
        results.append(item)
    return results


@router.post("/bulk-upload")
def bulk_upload_contacts(file: UploadFile = File(...), db: Session = Depends(get_db)):
    import io
    import csv

    ext = file.filename.rsplit(".", 1)[-1].lower() if file.filename else ""
    if ext not in ("csv", "xlsx", "xls"):
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type: .{ext}. Use .csv, .xlsx, or .xls",
        )

    content = file.file.read()
    rows = []

    if ext == "csv":
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        for row in reader:
            rows.append(row)
    else:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        ws = wb.active
        headers = [str(c.value).strip().lower().replace(" ", "_") for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(v is not None for v in row):
                rows.append(dict(zip(headers, [str(v) if v is not None else "" for v in row])))

    column_map = {
        "first_name": "first_name", "firstname": "first_name", "first name": "first_name",
        "last_name": "last_name", "lastname": "last_name", "last name": "last_name",
        "email": "email", "e-mail": "email",
        "designation": "designation", "title": "designation", "job_title": "designation", "job title": "designation", "role": "designation",
        "phone": "phone", "telephone": "phone", "mobile": "phone",
        "linkedin_url": "linkedin_url", "linkedin": "linkedin_url", "linkedin url": "linkedin_url", "linkedin profile": "linkedin_url",
        "company": "company_name", "company_name": "company_name", "organization": "company_name",
        "tags": "tags",
        "notes": "notes", "note": "notes",
        "source": "source",
        "status": "status",
        "lead_score": "lead_score", "score": "lead_score",
    }

    created = 0
    skipped = 0
    errors = []

    for i, row in enumerate(rows):
        try:
            mapped = {}
            for k, v in row.items():
                target = column_map.get(k.strip().lower(), None)
                if target:
                    mapped[target] = v.strip() if v else ""

            first = mapped.get("first_name", "")
            last = mapped.get("last_name", "")
            if not first and not last:
                skipped += 1
                continue
            if not first:
                first = last
                last = ""

            company_name = mapped.get("company_name", "")
            company_id = None
            if company_name:
                existing = db.query(Company).filter(Company.name.ilike(f"%{company_name}%")).first()
                if existing:
                    company_id = existing.id
                else:
                    company = Company(name=company_name)
                    db.add(company)
                    db.flush()
                    company_id = company.id

            email = mapped.get("email")

            source_str = mapped.get("source", "manual").lower()
            try:
                source = ContactSource(source_str)
            except ValueError:
                source = ContactSource.MANUAL

            status_str = mapped.get("status", "lead").lower()
            try:
                status = ContactStatus(status_str)
            except ValueError:
                status = ContactStatus.LEAD

            score_str = mapped.get("lead_score", "0")
            try:
                score = int(score_str)
            except (ValueError, TypeError):
                score = 0

            contact = Contact(
                first_name=first,
                last_name=last,
                company_id=company_id,
                designation=mapped.get("designation"),
                email=email,
                phone=mapped.get("phone"),
                linkedin_url=mapped.get("linkedin_url"),
                source=source,
                status=status,
                lead_score=score,
                tags=mapped.get("tags"),
                notes=mapped.get("notes"),
            )
            db.add(contact)
            db.flush()

            score_contact(db, contact.id)
            index_contact(
                contact_id=contact.id,
                first_name=contact.first_name,
                last_name=contact.last_name,
                email=contact.email,
                designation=contact.designation,
                tags=contact.tags,
                notes=contact.notes,
            )
            new_lead_webhook(
                contact_id=str(contact.id),
                contact_name=contact.full_name,
                email=email or "",
            )
            created += 1
        except Exception as e:
            errors.append({"row": i + 2, "error": str(e)})

    db.commit()
    return {
        "created": created,
        "skipped": skipped,
        "errors": errors,
        "total_rows": len(rows),
    }


@router.get("/search")
def search_contacts(
    q: str = Query(..., min_length=1),
    limit: int = Query(10, ge=1, le=50),
    db: Session = Depends(get_db),
):
    return search(query=q, limit=limit, db=db, type_filter="contact")


@router.get("/{contact_id}", response_model=ContactResponse)
def get_contact(contact_id: str, db: Session = Depends(get_db)):
    contact = db.query(Contact).filter(Contact.id == contact_id).first()
    if not contact:
        raise HTTPException(status_code=404, detail="Contact not found")
    return contact


@router.post("", response_model=ContactResponse, status_code=201)
def create_contact(body: ContactCreate, db: Session = Depends(get_db)):
    contact = Contact(
        first_name=body.first_name,
        last_name=body.last_name,
        company_id=(
            uuid.UUID(body.company_id) if body.company_id else None
        ),
        designation=body.designation,
        email=body.email,
        phone=body.phone,
        linkedin_url=body.linkedin_url,
        twitter_url=body.twitter_url,
        source=body.source,
        status=body.status,
        tags=body.tags,
        notes=body.notes,
    )
    db.add(contact)
    db.commit()
    db.refresh(contact)

    score_contact(db, contact.id)
    index_contact(
        contact_id=contact.id,
        first_name=contact.first_name,
        last_name=contact.last_name,
        email=contact.email,
        designation=contact.designation,
        tags=contact.tags,
        notes=contact.notes,
    )
    new_lead_webhook(
        contact_id=str(contact.id),
        contact_name=contact.full_name,
        email=contact.email or "",
    )

    return contact


@router.put("/{contact_id}", response_model=ContactResponse)
def update_contact(
    contact_id: str,
    body: ContactUpdate,
    db: Session = Depends(get_db),
):
    contact = db.query(Contact).filter(Contact.id == contact_id).first()
    if not contact:
        raise HTTPException(status_code=404, detail="Contact not found")

    update_data = body.model_dump(exclude_unset=True)
    for key, value in update_data.items():
        if key == "company_id" and value is not None:
            setattr(contact, key, uuid.UUID(value))
        elif value is not None:
            setattr(contact, key, value)

    db.commit()
    db.refresh(contact)
    return contact


@router.delete("/{contact_id}", status_code=204)
def delete_contact(contact_id: str, db: Session = Depends(get_db)):
    contact = db.query(Contact).filter(Contact.id == contact_id).first()
    if not contact:
        raise HTTPException(status_code=404, detail="Contact not found")
    delete_index(contact.id)
    db.delete(contact)
    db.commit()
